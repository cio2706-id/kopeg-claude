#!/usr/bin/env python3
"""Analyze loan Excel files to understand structure and installment calculations."""

import sys
import os

base_dir = "/home/user/kopeg-claude/Pinjaman Anggota/kartu pinjaman/"

# ============================================================
# FILE 1: PERHITUNGAN PINJAMAN BARANG.xlsx (openpyxl for .xlsx)
# ============================================================
print("=" * 80)
print("FILE 1: PERHITUNGAN PINJAMAN BARANG.xlsx")
print("=" * 80)

try:
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(base_dir, "PERHITUNGAN PINJAMAN BARANG.xlsx"), data_only=False)
    print(f"Sheet names: {wb.sheetnames}")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n--- Sheet: '{sheet_name}' ---")
        print(f"Dimensions: {ws.dimensions}")
        print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")

        # Print all rows (it's likely not huge)
        print("\nFull content (values):")
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 60), values_only=False):
            row_data = []
            for cell in row:
                if cell.value is not None:
                    row_data.append(f"{cell.coordinate}={cell.value}")
            if row_data:
                print("  " + " | ".join(row_data))

        # Now show formulas
        print("\nFormulas in this sheet:")
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for cell in row:
                if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                    print(f"  {cell.coordinate}: {cell.value}")

    # Also load with data_only=True to see computed values
    wb2 = openpyxl.load_workbook(os.path.join(base_dir, "PERHITUNGAN PINJAMAN BARANG.xlsx"), data_only=True)
    for sheet_name in wb2.sheetnames:
        ws2 = wb2[sheet_name]
        print(f"\n--- Sheet '{sheet_name}' computed values ---")
        for row in ws2.iter_rows(min_row=1, max_row=min(ws2.max_row, 60), values_only=False):
            row_data = []
            for cell in row:
                if cell.value is not None:
                    row_data.append(f"{cell.coordinate}={cell.value}")
            if row_data:
                print("  " + " | ".join(row_data))

except Exception as e:
    print(f"Error reading PERHITUNGAN PINJAMAN BARANG.xlsx: {e}")
    import traceback; traceback.print_exc()

# ============================================================
# FILE 2: master pinjaman Reguler.xls (xlrd for .xls)
# ============================================================
print("\n" + "=" * 80)
print("FILE 2: master pinjaman Reguler.xls")
print("=" * 80)

try:
    import xlrd
    wb = xlrd.open_workbook(os.path.join(base_dir, "master pinjaman Reguler.xls"), formatting_info=True)
    print(f"Sheet names: {wb.sheet_names()}")

    for idx, sheet_name in enumerate(wb.sheet_names()):
        ws = wb.sheet_by_index(idx)
        print(f"\n--- Sheet: '{sheet_name}' ---")
        print(f"Rows: {ws.nrows}, Cols: {ws.ncols}")

        # Print first 50 rows to understand structure
        print(f"\nFirst {min(ws.nrows, 50)} rows:")
        for r in range(min(ws.nrows, 50)):
            row_data = []
            for c in range(ws.ncols):
                cell = ws.cell(r, c)
                if cell.value not in (None, '', 0, 0.0):
                    row_data.append(f"[{r},{c}]({cell.ctype})={cell.value}")
            if row_data:
                print(f"  Row {r}: " + " | ".join(row_data))

        # If more rows, show structure summary
        if ws.nrows > 50:
            print(f"\n  ... ({ws.nrows - 50} more rows)")
            # Show a few rows from middle and end
            for r in [ws.nrows // 2, ws.nrows - 2, ws.nrows - 1]:
                row_data = []
                for c in range(ws.ncols):
                    cell = ws.cell(r, c)
                    if cell.value not in (None, '', 0, 0.0):
                        row_data.append(f"[{r},{c}]({cell.ctype})={cell.value}")
                if row_data:
                    print(f"  Row {r}: " + " | ".join(row_data))

except Exception as e:
    print(f"Error reading master pinjaman Reguler.xls: {e}")
    import traceback; traceback.print_exc()

# ============================================================
# FILE 3: master pinjaman Khusus.xls (xlrd for .xls)
# ============================================================
print("\n" + "=" * 80)
print("FILE 3: master pinjaman Khusus.xls")
print("=" * 80)

try:
    import xlrd
    wb = xlrd.open_workbook(os.path.join(base_dir, "master pinjaman Khusus.xls"), formatting_info=True)
    print(f"Sheet names: {wb.sheet_names()}")

    for idx, sheet_name in enumerate(wb.sheet_names()):
        ws = wb.sheet_by_index(idx)
        print(f"\n--- Sheet: '{sheet_name}' ---")
        print(f"Rows: {ws.nrows}, Cols: {ws.ncols}")

        # Print first 50 rows
        print(f"\nFirst {min(ws.nrows, 50)} rows:")
        for r in range(min(ws.nrows, 50)):
            row_data = []
            for c in range(ws.ncols):
                cell = ws.cell(r, c)
                if cell.value not in (None, '', 0, 0.0):
                    row_data.append(f"[{r},{c}]({cell.ctype})={cell.value}")
            if row_data:
                print(f"  Row {r}: " + " | ".join(row_data))

        if ws.nrows > 50:
            print(f"\n  ... ({ws.nrows - 50} more rows)")
            for r in [ws.nrows // 2, ws.nrows - 2, ws.nrows - 1]:
                row_data = []
                for c in range(ws.ncols):
                    cell = ws.cell(r, c)
                    if cell.value not in (None, '', 0, 0.0):
                        row_data.append(f"[{r},{c}]({cell.ctype})={cell.value}")
                if row_data:
                    print(f"  Row {r}: " + " | ".join(row_data))

except Exception as e:
    print(f"Error reading master pinjaman Khusus.xls: {e}")
    import traceback; traceback.print_exc()
